"""Export solved CANOE/TEMOA output tables to an Excel workbook.

This script reads ``Output*`` tables from a selected SQLite model database and
writes each available table to a separate worksheet in a single Excel workbook
for rapid inspection and analysis.

Two execution workflows are supported:

1. Interactive CLI selection of a model run and database from ``output_files/``.
2. Programmatic invocation from ``main_run.py`` after a successful model solve.

The export workflow validates worksheet dimensions against Excel limits,
generates valid and unique worksheet names, and applies basic formatting to
improve workbook readability.
"""

from __future__ import annotations

import argparse
import re
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter


# =============================================================================
# Project paths
# =============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = PROJECT_ROOT / "output_files"


# =============================================================================
# Default output table names
# =============================================================================

DEFAULT_OUTPUT_TABLES = [
    "OutputBuiltCapacity",
    "OutputCost",
    "OutputCurtailment",
    "OutputDualVariable",
    "OutputEmission",
    "OutputFlowIn",
    "OutputFlowOut",
    "OutputNetCapacity",
    "OutputObjective",
    "OutputRetiredCapacity",
    "OutputStorageLevel",
    "OutputFlowOutSummary",
]

DEFAULT_WORKBOOK_NAME = "output_tables.xlsx"
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384


# =============================================================================
# Interactive run selection helpers
# =============================================================================

def find_output_runs(output_root: Path = OUTPUT_ROOT) -> list[Path]:
    """Find model-run directories containing SQLite databases.

    The output root is expected to contain timestamped model-run directories. A
    directory is considered a valid run when it contains at least one ``.sqlite``
    file. Valid run directories are returned in sorted order.

    Parameters
    ----------
    output_root : Path, default=OUTPUT_ROOT
        Root directory containing timestamped CANOE/TEMOA model-run outputs.

    Returns
    -------
    list[Path]
        Sorted paths to model-run directories containing at least one SQLite
        database.

    Raises
    ------
    FileNotFoundError
        If ``output_root`` does not exist or no valid model-run directories are
        found.
    """

    if not output_root.exists():
        raise FileNotFoundError(f"Output root not found: {output_root}")

    runs = sorted(
        [
            path
            for path in output_root.iterdir()
            if path.is_dir() and any(path.glob("*.sqlite"))
        ]
    )

    if not runs:
        raise FileNotFoundError(
            f"No output runs with SQLite files found in {output_root}"
        )

    return runs


def select_from_list(options: list[Path], label: str) -> Path:
    """Prompt the user to select a path from a numbered list.

    The available paths are printed using their filenames and zero-based indices.
    The user is repeatedly prompted until they enter an integer corresponding to a
    valid option.

    Parameters
    ----------
    options : list[Path]
        Paths available for selection.
    label : str
        Human-readable label used in prompts and error messages.

    Returns
    -------
    Path
        Path selected by the user.

    Raises
    ------
    FileNotFoundError
        If ``options`` is empty.
    """

    if not options:
        raise FileNotFoundError(f"No {label} options found.")

    print(f"\nAvailable {label}:")
    for index, path in enumerate(options):
        print(f"  [{index}] {path.name}")

    while True:
        choice = input(f"\nSelect {label} index: ").strip()

        try:
            index = int(choice)
        except ValueError:
            print("Please enter an integer.")
            continue

        if 0 <= index < len(options):
            return options[index]

        print("Invalid selection.")


def find_sqlite_files(run_dir: Path) -> list[Path]:
    """Find SQLite databases in a selected model-run directory.

    The selected run directory is searched for files with the ``.sqlite`` suffix.
    Matching database paths are returned in sorted order.

    Parameters
    ----------
    run_dir : Path
        Model-run directory to search for SQLite database files.

    Returns
    -------
    list[Path]
        Sorted paths to SQLite databases found in ``run_dir``.

    Raises
    ------
    FileNotFoundError
        If no SQLite databases are found in ``run_dir``.
    """

    sqlite_files = sorted(run_dir.glob("*.sqlite"))

    if not sqlite_files:
        raise FileNotFoundError(f"No SQLite files found in {run_dir}")

    return sqlite_files


def select_run_database(output_root: Path = OUTPUT_ROOT) -> tuple[Path, Path]:
    """Select a model run and its SQLite database through CLI prompts.

    Available model-run directories are discovered under ``output_root`` and
    presented for numbered selection. The selected run is then searched for SQLite
    databases. When exactly one database begins with ``"solved_"``, it is selected
    automatically; otherwise, the user is prompted to choose from all databases in
    the run directory.

    Parameters
    ----------
    output_root : Path, default=OUTPUT_ROOT
        Root directory containing timestamped CANOE/TEMOA model-run outputs.

    Returns
    -------
    tuple[Path, Path]
        Selected model-run directory and selected SQLite database path.

    Raises
    ------
    FileNotFoundError
        If the output root does not exist, no valid model runs are found, or the
        selected run contains no SQLite databases.
    """

    runs = find_output_runs(output_root)
    selected_run = select_from_list(runs, "model runs")

    sqlite_files = find_sqlite_files(selected_run)

    solved_files = [
        path
        for path in sqlite_files
        if path.name.startswith("solved_")
    ]

    if len(solved_files) == 1:
        selected_db = solved_files[0]
        print(f"\nSelected solved database: {selected_db.name}")
    else:
        selected_db = select_from_list(sqlite_files, "SQLite databases")

    return selected_run, selected_db


# =============================================================================
# SQLite export helpers
# =============================================================================

def list_existing_output_tables(db_path: Path) -> list[str]:
    """Return the names of output tables present in a SQLite database.

    The database schema is queried for tables whose names begin with
    ``"Output"``. Matching table names are returned in alphabetical order.

    Parameters
    ----------
    db_path : Path
        Path to the SQLite database to inspect.

    Returns
    -------
    list[str]
        Alphabetically sorted names of all tables matching ``Output*``.
    """

    query = """
        SELECT name
        FROM sqlite_master
        WHERE type = 'table'
          AND name LIKE 'Output%'
        ORDER BY name
    """

    with sqlite3.connect(db_path) as connection:
        tables = pd.read_sql_query(query, connection)["name"].tolist()

    return tables


def make_excel_sheet_name(table_name: str, used_sheet_names: set[str]) -> str:
    """Return a valid and unique Excel worksheet name for a table.

    Characters prohibited in Excel worksheet names are replaced with underscores,
    leading and trailing whitespace is removed, and the result is limited to
    Excel's 31-character worksheet-name limit. If the normalized name has already
    been used, an incrementing numeric suffix is appended until a unique name is
    found. The selected name is added to ``used_sheet_names`` before being
    returned.

    Parameters
    ----------
    table_name : str
        SQLite table name to convert into an Excel-compatible worksheet name.
    used_sheet_names : set[str]
        Worksheet names already assigned within the current workbook. This set is
        updated in place with the generated name.

    Returns
    -------
    str
        Valid worksheet name that is unique within ``used_sheet_names``.
    """

    sheet_name = re.sub(r"[\\/*?:\[\]]", "_", table_name).strip() or "Sheet"
    sheet_name = sheet_name[:31]

    if sheet_name not in used_sheet_names:
        used_sheet_names.add(sheet_name)
        return sheet_name

    counter = 2
    while True:
        suffix = f"_{counter}"
        candidate = f"{sheet_name[:31 - len(suffix)]}{suffix}"

        if candidate not in used_sheet_names:
            used_sheet_names.add(candidate)
            return candidate

        counter += 1


def validate_excel_sheet_size(table: str, df: pd.DataFrame) -> None:
    """Validate that a table fits within a single Excel worksheet.

    The table dimensions are checked against Excel's maximum worksheet size. The
    row validation includes the header row written by pandas. A ``ValueError`` is
    raised when either the row or column limit would be exceeded.

    Parameters
    ----------
    table : str
        Name of the source SQLite table, used in validation error messages.
    df : pd.DataFrame
        Table data intended for export to one Excel worksheet.

    Returns
    -------
    None

    Raises
    ------
    ValueError
        If the DataFrame, including its header row, exceeds Excel's maximum row
        limit or contains more columns than Excel permits in one worksheet.
    """

    # Excel's row limit includes the header row written by pandas.
    if len(df) + 1 > EXCEL_MAX_ROWS:
        raise ValueError(
            f"{table} has {len(df):,} rows. This exceeds Excel's worksheet "
            f"limit of {EXCEL_MAX_ROWS:,} rows including the header. "
            "Use CSV for this table or filter/split it before export."
        )

    if len(df.columns) > EXCEL_MAX_COLUMNS:
        raise ValueError(
            f"{table} has {len(df.columns):,} columns. This exceeds Excel's "
            f"worksheet limit of {EXCEL_MAX_COLUMNS:,} columns."
        )


def format_excel_sheet(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    """Apply basic formatting to an exported Excel worksheet.

    The worksheet header row is frozen, an autofilter is applied across the populated
    range, and each column width is adjusted from the longest non-null value or
    column heading. Widths are constrained to a minimum of 10 and a maximum of 40
    characters to preserve readability without creating excessively wide sheets.

    Parameters
    ----------
    writer : pd.ExcelWriter
        Active Excel writer containing the worksheet to format.
    sheet_name : str
        Name of the worksheet associated with ``df``.
    df : pd.DataFrame
        Exported table used to determine filter bounds and column widths.

    Returns
    -------
    None
    """

    worksheet = writer.sheets[sheet_name]
    worksheet.freeze_panes = "A2"

    if len(df.columns) > 0:
        worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, column_name in enumerate(df.columns, start=1):
        values = df[column_name].dropna().astype(str)
        max_value_width = values.str.len().max() if not values.empty else 0
        width = max(len(str(column_name)), int(max_value_width)) + 2
        width = min(max(width, 10), 40)

        column_letter = get_column_letter(column_index)
        worksheet.column_dimensions[column_letter].width = width


def export_output_tables(
    db_path: Path,
    output_dir: Path,
    tables: list[str] | None = None,
    skip_empty: bool = True,
) -> list[Path]:
    """Export selected CANOE/TEMOA output tables to one Excel workbook.

    The SQLite database is inspected for tables whose names begin with ``"Output"``.
    When ``tables`` is omitted, all discovered output tables are considered for
    export. Otherwise, only the requested table names are processed. Missing tables
    are skipped with a console message, and empty tables are skipped when
    ``skip_empty`` is true.

    Each exported table is written to a separate, valid, uniquely named worksheet
    within ``output_tables.xlsx``. Worksheet dimensions are validated against Excel
    limits, and basic formatting is applied after writing.

    The function preserves the public interface used by ``main_run.py`` by accepting
    an output directory and returning a list of generated paths. The returned list
    contains the single Excel workbook path.

    Parameters
    ----------
    db_path : Path
        Path to the solved CANOE/TEMOA SQLite database.
    output_dir : Path
        Directory where the Excel workbook will be written. The directory is
        created when it does not already exist.
    tables : list[str] | None, default=None
        Output-table names to export. When omitted, all existing ``Output*`` tables
        in the database are selected.
    skip_empty : bool, default=True
        Whether to skip selected tables containing no rows.

    Returns
    -------
    list[Path]
        List containing the path to the generated Excel workbook.

    Raises
    ------
    FileNotFoundError
        If ``db_path`` does not exist.
    ValueError
        If the database contains no ``Output*`` tables, no selected tables are
        exported, or an exported table exceeds Excel worksheet limits.
    """

    db_path = Path(db_path)
    output_dir = Path(output_dir)

    if not db_path.exists():
        raise FileNotFoundError(f"SQLite database not found: {db_path}")

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / DEFAULT_WORKBOOK_NAME

    existing_tables = list_existing_output_tables(db_path)

    if not existing_tables:
        raise ValueError(f"No Output* tables found in database: {db_path}")

    selected_tables = tables if tables is not None else existing_tables

    exported_paths: list[Path] = []
    used_sheet_names: set[str] = set()
    written_tables = 0

    with sqlite3.connect(db_path) as connection:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for table in selected_tables:
                if table not in existing_tables:
                    print(f"[Skip] Table not found: {table}")
                    continue

                df = pd.read_sql_query(f'SELECT * FROM "{table}"', connection)

                if df.empty and skip_empty:
                    print(f"[Skip] Empty table: {table}")
                    continue

                validate_excel_sheet_size(table, df)

                sheet_name = make_excel_sheet_name(table, used_sheet_names)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                format_excel_sheet(writer, sheet_name, df)
                written_tables += 1

                print(
                    f"[Exported] {table}: {len(df):,} row(s) "
                    f"→ {output_path.name} [{sheet_name}]"
                )

    if written_tables == 0:
        raise ValueError("No Output* tables were exported. All selected tables were empty or missing.")

    exported_paths.append(output_path)
    return exported_paths


# =============================================================================
# CLI
# =============================================================================

def parse_args() -> argparse.Namespace:
    """Parse command-line arguments for exporting solved output tables to Excel.

    The command-line interface supports either interactive model-run selection or
    direct selection of a solved SQLite database. Users may override the output
    directory, restrict the export to specific ``Output*`` tables, and choose
    whether empty tables should be included in the workbook.

    Returns
    -------
    argparse.Namespace
        Parsed command-line arguments containing the output-root path, optional
        database path, optional workbook output directory, optional table-name
        selection, and empty-table inclusion flag.
    """

    parser = argparse.ArgumentParser(
        description="Export solved CANOE/TEMOA Output* SQLite tables to Excel."
    )

    parser.add_argument(
        "--output-root",
        type=Path,
        default=OUTPUT_ROOT,
        help="Root folder containing timestamped model-run output folders.",
    )

    parser.add_argument(
        "--db",
        type=Path,
        default=None,
        help=(
            "Optional direct path to a solved SQLite database. "
            "If omitted, choose a run interactively."
        ),
    )

    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help=(
            "Optional Excel output directory. "
            "Defaults to <selected_run>/excel_outputs in interactive mode, "
            "or <db_parent>/excel_outputs when --db is provided."
        ),
    )

    parser.add_argument(
        "--tables",
        nargs="*",
        default=None,
        help=(
            "Optional list of Output* tables to export. "
            "Defaults to all existing Output* tables."
        ),
    )

    parser.add_argument(
        "--include-empty",
        action="store_true",
        help="Write empty Output* tables instead of skipping them.",
    )

    return parser.parse_args()


def main() -> None:
    """Run the command-line workflow for exporting output tables to Excel.

    Command-line arguments are parsed to determine whether the SQLite database is
    provided directly or selected interactively from the configured output root.
    The workbook output directory is then resolved from ``--out`` or derived from
    the selected run or database directory.

    The selected ``Output*`` tables are exported using ``export_output_tables``,
    with empty-table handling controlled by the ``--include-empty`` option. The
    resolved configuration and final workbook location are reported to the console.

    Returns
    -------
    None
    """

    args = parse_args()

    if args.db is not None:
        db_path = args.db
        output_dir = args.out or db_path.parent / "excel_outputs"
    else:
        selected_run, db_path = select_run_database(args.output_root)
        output_dir = args.out or selected_run / "excel_outputs"

    print("\nExcel export configuration")
    print("--------------------------")
    print(f"Database: {db_path}")
    print(f"Output:   {output_dir}")

    exported_paths = export_output_tables(
        db_path=db_path,
        output_dir=output_dir,
        tables=args.tables,
        skip_empty=not args.include_empty,
    )

    print("\nExcel export complete.")
    print(f"Exported {len(exported_paths):,} workbook(s) to: {output_dir}")


if __name__ == "__main__":
    main()
